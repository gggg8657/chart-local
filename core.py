"""chart local — 표 읽기·열 타입 추론·그래프 후보 추천·차트용 데이터 준비 (stdlib 만)

표(table) = {"columns": [{"name", "type": num|date|cat, "unit"}], "rows": [[문자열…]…], "title": 제목 후보}
그래프 명세(spec) = {"kind", "x", "ys": [...], "group", "agg", 스타일…} — 렌더러(svg.py)·내보내기(export.py) 공통 입력.
"""
import csv
import io
import math
import re
import fin

MAX_ROWS = 5000
MAX_COLS = 60

# ── 숫자·단위 ──────────────────────────────────────────────────────────────
_NUM = re.compile(r"^([+\-−–]?)\s*([₩$€¥£]?)\s*(\d{1,3}(?:,\d{3})+|\d+)?(\.\d+)?(?:[eE]([+\-]?\d+))?\s*(.*?)$")
_UNIT_OK = re.compile(r"^(%|‰|[A-Za-zµμ°Ω/²³·.\-]{1,10}|[가-힣]{1,4}|[A-Za-z]+/[A-Za-z가-힣]+|천?[만억조]?원|[천만억조]+[가-힣]{0,2})$")
_EMPTY = {"", "-", "–", "—", "n/a", "N/A", "NA", "na", "null", "None", "nan", "NaN", "#N/A", "#DIV/0!", "."}


def split_num(s):
    """'1,234.5억원' → (1234.5, '억원'); 숫자가 아니면 (None, None). 괄호 음수 (123) 지원."""
    if s is None:
        return None, None
    if isinstance(s, (int, float)) and not isinstance(s, bool):
        return (float(s), "") if math.isfinite(s) else (None, None)
    t = str(s).strip().replace(" ", " ")
    if t in _EMPTY:
        return None, None
    neg = False
    if len(t) > 2 and t[0] == "(" and t[-1] == ")":
        neg, t = True, t[1:-1].strip()
    if t.startswith("△") or t.startswith("▽"):  # 공문서 음수 표기
        neg, t = True, t[1:].strip()
    m = _NUM.match(t)
    if not m or not (m.group(3) or m.group(4)):
        return None, None
    sign, _cur, ip, fp, ex, rest = m.groups()
    rest = rest.strip()
    if rest and not _UNIT_OK.match(rest):
        return None, None
    v = float((ip or "0").replace(",", "") + (fp or "") + (("e" + ex) if ex else ""))
    if sign in "-−–" and sign:
        v = -v
    if neg:
        v = -v
    return v, rest


def num(s):
    return split_num(s)[0]


# ── 날짜·기간 ──────────────────────────────────────────────────────────────
_DATE_PATS = [
    re.compile(r"^(19|20)\d{2}$"),
    re.compile(r"^(19|20)\d{2}\s*년(도)?$"),
    re.compile(r"^(19|20)\d{2}[-./]\s?\d{1,2}([-./]\s?\d{1,2})?\.?$"),
    re.compile(r"^(19|20)\d{2}\s*년\s*\d{1,2}\s*월(\s*\d{1,2}\s*일)?$"),
    re.compile(r"^\d{1,2}\s*월$"),
    re.compile(r"^(19|20)?\d{2}\s*[-./]?\s*[1-4]\s*(Q|분기)$", re.I),
    re.compile(r"^((19|20)\d{2}\s*)?[1-4]\s*분기$"),
    re.compile(r"^((19|20)\d{2}\s*)?Q[1-4]$", re.I),
    re.compile(r"^((19|20)\d{2}\s*)?(상|하)반기$"),
    re.compile(r"^(19|20)\d{2}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2})?$"),
]
_TIME_NAME = re.compile(r"(연도|년도|년|월|일자|날짜|기간|분기|반기|시점|date|year|month|quarter|time)", re.I)


def is_date(s):
    t = str(s).strip()
    return any(p.match(t) for p in _DATE_PATS)


_PERIOD_MORE = re.compile(r"^\d{1,2}\s*(주차|일)$")
_CAT_HEAD = re.compile(r"(범주|항목|구분|분류|기간|월|분기|연도|년도|주차|요일|category|item|period|month|quarter)", re.I)


def period_column(name, vals):
    """'1월…12월'·'1분기'·'Q1'·'2026년'·'1주차' 처럼 숫자+기간 글자인 값만 있는 열 → 순서 있는 기간(날짜) 열로 볼지.
    split_num 은 '1월' 을 숫자 1(단위 월)로 읽으므로 따로 본다. '3개'·'5명' 같은 진짜 수량은 해당 없음.
    조건: 비지 않은 값이 모두 기간 꼴이고 숫자만 있는 값은 없음 + (값이 서로 다름 또는 머리글이 범주·기간 말)"""
    vs = [str(v).strip() for v in vals]
    if not vs or any(re.fullmatch(r"[\d.,\s]+", v) for v in vs):
        return False
    if not all(is_date(v) or _PERIOD_MORE.match(v) for v in vs):
        return False
    head = bool(_CAT_HEAD.search(name or ""))
    if all(_PERIOD_MORE.match(v) for v in vs):  # '3일'·'2주차' 는 기간이 아닐 수도 → 1씩 늘어나는 차례이거나 머리글이 기간 말일 때만
        ns = [int(re.match(r"\d+", v).group()) for v in vs]
        return head or (len(ns) >= 3 and all(b == a + 1 for a, b in zip(ns, ns[1:])))
    return head or len(set(vs)) == len(vs)


def date_key(s):
    """정렬용 키 — 숫자만 뽑아 튜플로. 분기·반기는 그 순서대로."""
    t = str(s)
    nums = [int(x) for x in re.findall(r"\d+", t)]
    if "하반기" in t:
        nums.append(2)
    elif "상반기" in t:
        nums.append(1)
    return tuple(nums)


# ── 표 읽기 ──────────────────────────────────────────────────────────────
def read_text(text):
    """붙여 넣은 텍스트(엑셀 탭 구분·CSV·세미콜론) → 2차원 문자열 배열"""
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n").strip("\n")
    if not text.strip():
        return []
    lines = [l for l in text.split("\n")]
    if any("\t" in l for l in lines[:20]):
        delim = "\t"
    else:
        sample = "\n".join(lines[:20])
        counts = {d: min((l.count(d) for l in lines[:20] if l.strip()), default=0) for d in (",", ";", "|")}
        delim = max(counts, key=counts.get) if max(counts.values()) > 0 else None
        if delim == "," and re.search(r"\d,\d{3}", sample) and "\"" not in sample and counts.get(";", 0) == 0:
            # 천단위 쉼표 숫자가 섞인 '쉼표 구분' — csv 모듈이 따옴표 없는 1,234 를 쪼갠다: 공백 2개 이상 구분 시도
            if all(re.search(r"\S\s{2,}\S", l) for l in lines[:5] if l.strip()):
                delim = "  "
    if delim is None:  # 공백 2칸 이상 정렬된 표 또는 한 열
        return [re.split(r"\s{2,}", l.strip()) for l in lines if l.strip()]
    if delim == "  ":
        return [re.split(r"\s{2,}", l.strip()) for l in lines if l.strip()]
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def xlsx_sheets(blob):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def xlsx_source_colors(blob):
    """디지타이저가 내보낸 XLSX 의 '보정' 시트에 적힌 원본 그림 계열 색 → {계열 이름: '#rrggbb'} (없으면 {}).
    줄 모양: ['계열 색', 이름, '#rrggbb'] — 다른 표에는 이 시트·줄이 없으니 빈 지도"""
    import openpyxl
    out = {}
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        if "보정" not in wb.sheetnames:
            return out
        for row in wb["보정"].iter_rows(values_only=True, max_row=400):
            if len(row) >= 3 and row[0] == "계열 색" and isinstance(row[1], str) and isinstance(row[2], str) \
                    and re.fullmatch(r"#[0-9a-fA-F]{6}", row[2].strip()) and len(out) < 24:
                out[row[1].strip()[:80]] = row[2].strip().lower()
    finally:
        wb.close()
    return out


def read_xlsx(blob, sheet=None):
    import datetime
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
        out = []
        for row in ws.iter_rows(values_only=True):
            r = []
            for v in row[:MAX_COLS]:
                if v is None:
                    r.append("")
                elif isinstance(v, datetime.datetime):
                    r.append(v.strftime("%Y-%m-%d") if not (v.hour or v.minute) else v.strftime("%Y-%m-%d %H:%M"))
                elif isinstance(v, datetime.date):
                    r.append(v.isoformat())
                elif isinstance(v, float):
                    r.append(repr(round(v, 10)).rstrip("0").rstrip(".") if v != int(v) or abs(v) > 1e15 else str(int(v)))
                else:
                    r.append(str(v))
            out.append(r)
            if len(out) > MAX_ROWS + 10:
                break
        return out, ws.title
    finally:
        wb.close()


def read_csv_bytes(blob):
    for enc in ("utf-8-sig", "cp949", "euc-kr", "latin-1"):
        try:
            return read_text(blob.decode(enc))
        except UnicodeDecodeError:
            continue
    return []


# ── 정규화: 머리글·단위 행·제목 행·전치 ────────────────────────────────────
_UNIT_CELL = re.compile(r"^[\(\[〔<]?\s*(단위\s*[:：]?\s*)?([^\d\s\(\)\[\]]{1,10})\s*[\)\]〕>]?$")
_HEAD_UNIT = re.compile(r"^(.*?)\s*[\(\[〔]\s*(?:단위\s*[:：]?\s*)?([^\(\)\[\]〔〕]{1,12})\s*[\)\]〕]\s*$")
_GLOBAL_UNIT = re.compile(r"[\(\[]?\s*단위\s*[:：]\s*([^\)\]\s,]{1,10})\s*[\)\]]?")


def _numlike(c):
    return num(c) is not None


def normalize(rows, transpose=None):
    """원시 2차원 배열 → table + notes(무엇을 자동 처리했는지 한국어 설명)"""
    notes = []
    rows = [[("" if c is None else str(c)).strip() for c in r] for r in rows]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        raise ValueError("표 데이터가 없습니다")
    width = min(MAX_COLS, max(len(r) for r in rows))
    rows = [(r + [""] * width)[:width] for r in rows]
    keep = [j for j in range(width) if any(r[j] for r in rows)]
    rows = [[r[j] for j in keep] for r in rows]
    width = len(keep)
    title, gunit = "", ""
    # 제목·단위 줄: 칸이 하나뿐인 앞쪽 줄
    while len(rows) > 2 and sum(1 for c in rows[0] if c) == 1 and sum(1 for c in rows[1] if c) >= 2 and width >= 2:
        only = next(c for c in rows[0] if c)
        gm = _GLOBAL_UNIT.search(only)
        if gm:
            gunit = gm.group(1)
            rest = _GLOBAL_UNIT.sub("", only).strip(" -·:")
            if rest and not title:
                title = rest
        elif not title:
            title = only
        rows = rows[1:]
        notes.append(f"첫 줄 '{only[:30]}' 을(를) {'단위' if gm else '제목'}로 사용")
    # 머리글 판단
    head = rows[0]
    body = rows[1:]
    header = True
    if body:
        num_cols = [j for j in range(width) if sum(_numlike(r[j]) for r in body) >= max(1, 0.7 * sum(1 for r in body if r[j]))]
        head_num = sum(1 for j in num_cols if _numlike(head[j]) and not is_date(head[j]))
        head_txt = sum(1 for c in head if c and not _numlike(c))
        if num_cols and head_num == len(num_cols) and head_txt == 0:
            header = False
    if not body:
        header = False
    if header:
        names = [c or f"열{j + 1}" for j, c in enumerate(head)]
        rows = body
    else:
        names = [f"열{j + 1}" for j in range(width)]
        notes.append("머리글이 없어 보여 열1·열2… 로 이름 붙임")
    # 머리글 바로 아래 단위 행
    units = [""] * width
    if len(rows) > 2 and all((not c) or (_UNIT_CELL.match(c) and not _numlike(c)) for c in rows[0]) \
            and any(rows[0]) and any(_numlike(c) for c in rows[1]):
        for j, c in enumerate(rows[0]):
            m = _UNIT_CELL.match(c) if c else None
            units[j] = m.group(2).strip() if m else ""
        rows = rows[1:]
        notes.append("머리글 아래 단위 행을 열 단위로 사용")
    # 머리글 안 단위: '예산(억원)'
    for j, n in enumerate(names):
        m = _HEAD_UNIT.match(n)
        if m and m.group(1).strip():
            names[j], units[j] = m.group(1).strip(), units[j] or m.group(2).strip()
    names = _dedupe(names)
    empty = [j for j in range(width) if rows and not any(str(r[j]).strip() not in _EMPTY for r in rows)]
    if empty and len(empty) < width:  # 머리글만 있고 값이 모두 빈 열(예: 디지타이저 '막대 끝(합계)') → 뺌
        notes.append("값이 모두 빈 열 " + ", ".join(f"'{names[j]}'" for j in empty) + " 을(를) 뺌")
        keep = [j for j in range(width) if j not in empty]
        rows = [[r[j] for j in keep] for r in rows]
        names, units = [names[j] for j in keep], [units[j] for j in keep]
        width = len(keep)
    if len(rows) > MAX_ROWS:
        notes.append(f"행이 많아 앞 {MAX_ROWS}행만 사용")
        rows = rows[:MAX_ROWS]
    table = {"columns": [{"name": n, "unit": u, "type": "cat"} for n, u in zip(names, units)], "rows": rows, "title": title}
    infer_types(table, gunit)
    # 연도가 머리글에 가로로 늘어선 표 → 전치 (첫 열이 항목 이름이고 나머지 머리글이 연도·기간)
    cols = table["columns"]
    wide = (header and len(cols) >= 4 and cols[0]["type"] == "cat" and all(is_date(c["name"]) for c in cols[1:])
            and all(c["type"] == "num" for c in cols[1:]) and len(rows) <= 12)
    if transpose is True or (transpose is None and wide):
        table = transpose_table(table)
        notes.append("기간이 열 머리글에 가로로 놓여 있어 행·열을 바꿈" if transpose is None else "행·열을 바꿈")
    return table, notes


def transpose_table(table):
    cols, rows = table["columns"], table["rows"]
    first = cols[0]
    names = [first["name"] or "구분"] + [r[0] or f"행{i + 1}" for i, r in enumerate(rows)]
    unit = next((c["unit"] for c in cols[1:] if c["unit"]), "")
    new_rows = [[c["name"]] + [r[j + 1] if j + 1 < len(r) else "" for r in rows] for j, c in enumerate(cols[1:])]
    t = {"columns": [{"name": n, "unit": "" if i == 0 else unit, "type": "cat"} for i, n in enumerate(_dedupe(names))],
         "rows": new_rows, "title": table.get("title", "")}
    infer_types(t)
    return t


def _dedupe(names):
    seen, out = {}, []
    for n in names:
        k = n
        while k in seen:
            seen[n] = seen.get(n, 1) + 1
            k = f"{n}_{seen[n]}"
        seen[k] = 1
        out.append(k)
    return out


def infer_types(table, gunit=""):
    """열 타입(num|date|cat)과 단위 추론. 값 뒤의 공통 접미사(%, kg, 억원 …)는 단위로 옮긴다."""
    rows = table["rows"]
    for j, c in enumerate(table["columns"]):
        vals = [r[j] for r in rows if j < len(r) and str(r[j]).strip() not in _EMPTY]
        if not vals:
            c["type"] = "cat"
            continue
        parsed = [split_num(v) for v in vals]
        nnum = sum(1 for v, _ in parsed if v is not None)
        dates = sum(1 for v in vals if is_date(v))
        sufs = {u for v, u in parsed if v is not None and u}
        if period_column(c["name"], vals):
            c["type"] = "date"
        elif dates >= 0.8 * len(vals) and not (nnum == len(vals) and not _TIME_NAME.search(c["name"]) and not all(re.fullmatch(r"(19|20)\d{2}", v.strip()) for v in vals)):
            c["type"] = "date"
        elif nnum >= 0.8 * len(vals):
            ints = [v for v, _ in parsed if v is not None]
            yearish = all(float(v).is_integer() and 1900 <= v <= 2100 for v in ints)
            if yearish and (_TIME_NAME.search(c["name"]) or (j == 0 and len(set(ints)) == len(ints))):
                c["type"] = "date"
            else:
                c["type"] = "num"
                if len(sufs) == 1 and not c.get("unit"):
                    c["unit"] = sufs.pop()
                if not c.get("unit") and gunit:
                    c["unit"] = gunit
        else:
            c["type"] = "cat"
    return table


def col_index(table, name):
    for j, c in enumerate(table["columns"]):
        if c["name"] == name:
            return j
    raise ValueError(f"열 '{name}' 이(가) 없습니다")


def col_values(table, name):
    j = col_index(table, name)
    return [r[j] if j < len(r) else "" for r in table["rows"]]


def col_nums(table, name):
    return [num(v) for v in col_values(table, name)]


def unit_of(table, name):
    try:
        return table["columns"][col_index(table, name)].get("unit", "")
    except ValueError:
        return ""


# ── 통계 ──────────────────────────────────────────────────────────────────
def quantile(sorted_vals, q):
    """엑셀 QUARTILE.INC 와 같은 선형 보간"""
    n = len(sorted_vals)
    if n == 1:
        return sorted_vals[0]
    pos = (n - 1) * q
    lo = int(math.floor(pos))
    hi = min(lo + 1, n - 1)
    return sorted_vals[lo] + (sorted_vals[hi] - sorted_vals[lo]) * (pos - lo)


def box_stats(vals):
    v = sorted(x for x in vals if x is not None)
    if not v:
        return None
    q1, med, q3 = quantile(v, .25), quantile(v, .5), quantile(v, .75)
    iqr = q3 - q1
    lo_f, hi_f = q1 - 1.5 * iqr, q3 + 1.5 * iqr
    inside = [x for x in v if lo_f <= x <= hi_f] or v
    return {"min": v[0], "q1": q1, "median": med, "q3": q3, "max": v[-1], "wlo": inside[0], "whi": inside[-1],
            "out": [x for x in v if x < lo_f or x > hi_f], "n": len(v), "mean": sum(v) / len(v)}


TREND_KINDS = {"col", "col_group", "bar", "bar_group", "col_stack", "bar_stack"}
STACK_TREND = {"col_stack", "bar_stack"}


def trend_mode(spec):
    """막대 추세선: 'linear'(최소제곱 직선) | 'ma'(이동평균, 창 trend_n) | None. 산점도의 trend(참/거짓)와는 따로"""
    t = spec.get("trend")
    return t if spec.get("kind") in TREND_KINDS and t in ("linear", "ma") else None


def trend_n(spec, n):
    try:
        w = int(spec.get("trend_n") or 3)
    except (TypeError, ValueError):
        w = 3
    return max(2, min(w, max(2, n - 1)))


def trend_values(vals, mode, w=3):
    """x = 1…n(엑셀 범주 축 추세선과 같음). linear: 빈 값을 뺀 점으로 최소제곱 직선을 모든 칸에. ma: 앞 w 개 단순 평균(앞 w−1 칸 빈칸)"""
    n = len(vals)
    if mode == "linear":
        lr = linreg(list(range(1, n + 1)), vals)
        if not lr:
            return [None] * n
        return [lr["slope"] * (i + 1) + lr["intercept"] for i in range(n)]
    out = [None] * n
    for i in range(w - 1, n):
        win = vals[i - w + 1:i + 1]
        if all(v is not None for v in win):
            out[i] = sum(win) / w
    return out


def bar_trends(spec, series):
    """보이는 계열 → 추세 계열 [{name, values, idx, src}] (누적 막대는 합계 하나). 추세선이 없으면 []"""
    mode = trend_mode(spec)
    if not mode or not series:
        return []
    n = len(series[0]["values"])
    w = trend_n(spec, n)
    if spec["kind"] in STACK_TREND:
        tot = [sum(s["values"][i] or 0 for s in series) if any(s["values"][i] is not None for s in series) else None for i in range(n)]
        return [{"name": "합계 추세", "values": trend_values(tot, mode, w), "idx": None, "src": "합계", "total": tot}]
    return [{"name": f"{s['name']} 추세", "values": trend_values(s["values"], mode, w), "idx": s["idx"], "src": s["name"]} for s in series]


def trend_label(spec, n=None):
    mode = trend_mode(spec)
    if not mode:
        return ""
    return "직선 추세(최소제곱)" if mode == "linear" else f"이동평균 {trend_n(spec, n or 99)} 추세"


def linreg(xs, ys):
    pts = [(x, y) for x, y in zip(xs, ys) if x is not None and y is not None]
    n = len(pts)
    if n < 3:
        return None
    mx = sum(p[0] for p in pts) / n
    my = sum(p[1] for p in pts) / n
    sxx = sum((p[0] - mx) ** 2 for p in pts)
    syy = sum((p[1] - my) ** 2 for p in pts)
    sxy = sum((p[0] - mx) * (p[1] - my) for p in pts)
    if sxx == 0:
        return None
    a = sxy / sxx
    b = my - a * mx
    r2 = (sxy * sxy) / (sxx * syy) if syy else 1.0
    return {"slope": a, "intercept": b, "r2": r2, "r": (sxy / math.sqrt(sxx * syy)) if syy else 1.0}


def nice_step(span, target=5):
    if span <= 0:
        return 1
    raw = span / target
    p = 10 ** math.floor(math.log10(raw))
    for m in (1, 2, 2.5, 5, 10):
        if raw <= m * p:
            return m * p
    return 10 * p


def hist_bins(vals, nbins=None):
    """깔끔한 구간 경계(1·2·2.5·5 배수) — [(lo, hi, count)]. 마지막 구간은 오른쪽 끝 포함."""
    v = sorted(x for x in vals if x is not None)
    if not v:
        return []
    lo, hi = v[0], v[-1]
    if lo == hi:
        return [(lo, hi, len(v))]
    k = nbins or max(5, min(20, int(math.ceil(math.log2(len(v)) + 1))))
    step = nice_step(hi - lo, k)
    start = math.floor(lo / step) * step
    edges = [start]
    while edges[-1] <= hi - 1e-12 * max(1, abs(hi)):
        edges.append(round(edges[-1] + step, 10))
    if len(edges) < 2:
        edges.append(start + step)
    out = []
    for i in range(len(edges) - 1):
        a, b = edges[i], edges[i + 1]
        last = i == len(edges) - 2
        out.append((a, b, sum(1 for x in v if a <= x < b or (last and x == b))))
    return out


# ── 서식 ──────────────────────────────────────────────────────────────────
def decimals_of(vals, cap=2):
    d = 0
    for v in vals:
        if v is None or not math.isfinite(v):
            continue
        s = f"{abs(v):.6f}".rstrip("0")
        frac = s.split(".")[1] if "." in s else ""
        d = max(d, len(frac))
    return min(d, cap)


def fmt(v, dec=None, unit=""):
    if v is None:
        return ""
    if dec is None:
        dec = decimals_of([v])
    a = abs(v)
    if a >= 1e12 and unit not in ("%",):
        s = f"{v / 1e12:,.{1 if a < 1e13 else 0}f}조"
    elif a >= 1e8 and unit not in ("%",):
        s = f"{v / 1e8:,.{1 if a < 1e9 else 0}f}억"
    else:
        s = f"{v:,.{dec}f}"
    if s.startswith("-"):
        s = "−" + s[1:]
    return s + ("%" if unit == "%" else "")


# ── 차트 데이터 준비 ───────────────────────────────────────────────────────
KINDS = {
    "col": "세로 막대", "bar": "가로 막대", "col_group": "묶은 세로 막대", "bar_group": "묶은 가로 막대",
    "col_stack": "누적 세로 막대", "col_stack100": "100% 누적 막대", "bar_stack": "누적 가로 막대",
    "line": "꺾은선", "area": "영역", "area_stack": "누적 영역", "scatter": "산점도", "pie": "원그래프",
    "doughnut": "도넛", "hist": "히스토그램", "box": "상자그림", "heatmap": "히트맵", "panels": "단위별 2단 비교",
    "index": "지수 비교(첫 값=100)", "dual": "이중축 (보조 y축)",
    "radar": "방사형", "sparkline": "스파크라인 (계열별 작은 추세선)", "candlestick": "캔들스틱 (시가·고가·저가·종가)", "ohlc": "OHLC 막대 (시가·고가·저가·종가)",
    "waterfall": "폭포 (증감 누적)", "funnel": "깔때기 (단계별 감소)", "bubble": "거품 (x·y·크기)", "step": "계단선",
    "lollipop": "막대사탕 (점+줄기)", "dumbbell": "덤벨 (두 값 비교)", "treemap": "트리맵", "bullet": "목표 대비 (불릿)",
    "rose": "방사 면적 (로즈)", "stack_line": "누적 막대 + 꺾은선", "semi_doughnut": "반원 도넛",
}
CAT_KINDS = {"dual", "col", "bar", "col_group", "bar_group", "col_stack", "col_stack100", "bar_stack", "line", "area",
             "area_stack", "pie", "doughnut", "heatmap", "panels", "index",
             "radar", "sparkline", "candlestick", "ohlc", "waterfall", "funnel", "step", "lollipop", "dumbbell", "treemap",
             "bullet", "rose", "stack_line", "semi_doughnut"}
NO_SORT = {"waterfall", "candlestick", "ohlc", "sparkline", "step"}
# 값 열(ys) 말고 따로 지정하는 열: 종류 → (명세 키, 계열 역할)
EXTRA_COL = {"dual": ("y2", "right"), "stack_line": ("y2", "line"), "bullet": ("target", "target"),
             "candlestick": ("volume", "volume"), "ohlc": ("volume", "volume")}


def visible_ys(spec):
    hidden = set(spec.get("hidden") or [])
    return [y for y in (spec.get("ys") or []) if y not in hidden]


def category_data(table, spec):
    """범주형 차트 공통: x 순서대로 cats + 계열별 값. group 이 있으면 긴 형식 → 넓은 형식 피벗.
    중복된 x 는 spec.agg(sum|mean) 로 묶는다. 반환: {cats, series:[{name, values, unit, idx}], agg_note}"""
    x = spec.get("x")
    ys = list(spec.get("ys") or [])
    kind = spec.get("kind")
    extra_key, role = EXTRA_COL.get(kind, (None, None))
    extra = spec.get(extra_key) if extra_key else None
    if extra:  # 이중축 y2(오른쪽 축)·누적+선 y2(선)·불릿 target(목표)·캔들 volume(거래량) — 값 열 뒤에 붙임
        ys = [y for y in ys if y != extra] + [extra]
    xs = col_values(table, x) if x else [str(i + 1) for i in range(len(table["rows"]))]
    agg = spec.get("agg") or "sum"
    order, seen = [], set()
    for v in xs:
        if v not in seen:
            seen.add(v)
            order.append(v)
    xtype = table["columns"][col_index(table, x)]["type"] if x else "cat"
    if xtype == "date":
        try:
            order.sort(key=date_key)
        except TypeError:
            pass
    note = ""
    series = []
    if spec.get("group"):
        y = ys[0]
        gs = col_values(table, spec["group"])
        vals = col_nums(table, y)
        levels = []
        for g in gs:
            if g not in levels:
                levels.append(g)
        acc = {}
        for xv, g, v in zip(xs, gs, vals):
            if v is None:
                continue
            acc.setdefault((xv, g), []).append(v)
        if any(len(a) > 1 for a in acc.values()):
            note = f"같은 ({x}, {spec['group']}) 조합이 여러 행이라 {'평균' if agg == 'mean' else '합계'}로 묶음"
        hidden = set(spec.get("hidden") or [])
        for i, g in enumerate(levels):
            vs = []
            for c in order:
                a = acc.get((c, g))
                vs.append(None if not a else (sum(a) / len(a) if agg == "mean" else sum(a)))
            series.append({"name": g or "(빈 값)", "values": vs, "unit": unit_of(table, y), "idx": i, "hidden": g in hidden})
    else:
        pos = {c: i for i, c in enumerate(order)}
        for i, y in enumerate(ys):
            vals = col_nums(table, y)
            acc = [[] for _ in order]
            for xv, v in zip(xs, vals):
                if v is not None:
                    acc[pos[xv]].append(v)
            if any(len(a) > 1 for a in acc):
                note = f"같은 '{x}' 값이 여러 행이라 {'평균' if agg == 'mean' else '합계'}로 묶음"
            vs = [None if not a else (sum(a) / len(a) if agg == "mean" else sum(a)) for a in acc]
            series.append({"name": y, "values": vs, "unit": unit_of(table, y), "idx": i, "right": y == extra and kind == "dual",
                           "role": role if (extra and y == extra) else None, "hidden": y in set(spec.get("hidden") or [])})
    cats = [c if c != "" else "(빈 값)" for c in order]
    # 정렬 (시간축은 정렬하지 않음)
    srt = spec.get("sort") or "none"
    if kind in NO_SORT:  # 순서 자체가 뜻인 종류(누적·시간 흐름)는 정렬하지 않음
        srt = "none"
    vis = [s for s in series if not s["hidden"] and not s.get("role")] or [s for s in series if not s["hidden"]]
    if srt in ("asc", "desc") and xtype != "date" and vis:
        tot = [sum((s["values"][i] or 0) for s in vis) for i in range(len(cats))]
        idx = sorted(range(len(cats)), key=lambda i: tot[i], reverse=srt == "desc")
        cats = [cats[i] for i in idx]
        for s in series:
            s["values"] = [s["values"][i] for i in idx]
    top = int(spec.get("top") or 0)
    if top and len(cats) > top:
        cats = cats[:top]
        for s in series:
            s["values"] = s["values"][:top]
    return {"cats": cats, "series": series, "note": note, "xtype": xtype}


def scatter_data(table, spec):
    x = spec["x"]
    xs = col_nums(table, x)
    out = []
    if spec.get("group"):
        gs = col_values(table, spec["group"])
        y = spec["ys"][0]
        ys = col_nums(table, y)
        levels = []
        for g in gs:
            if g not in levels:
                levels.append(g)
        for i, g in enumerate(levels):
            pts = [(a, b) for a, b, gg in zip(xs, ys, gs) if gg == g and a is not None and b is not None]
            out.append({"name": g, "points": pts, "idx": i, "hidden": g in set(spec.get("hidden") or [])})
    else:
        for i, y in enumerate(spec["ys"]):
            ys = col_nums(table, y)
            pts = [(a, b) for a, b in zip(xs, ys) if a is not None and b is not None]
            out.append({"name": y, "points": pts, "idx": i, "hidden": y in set(spec.get("hidden") or [])})
    return out


def dist_data(table, spec):
    """히스토그램·상자그림: 계열(열 또는 그룹)별 숫자 목록"""
    if spec.get("group"):
        y = spec["ys"][0]
        gs = col_values(table, spec["group"])
        vs = col_nums(table, y)
        levels = []
        for g in gs:
            if g not in levels:
                levels.append(g)
        return [{"name": g, "values": [v for v, gg in zip(vs, gs) if gg == g and v is not None], "idx": i,
                 "hidden": g in set(spec.get("hidden") or [])} for i, g in enumerate(levels)]
    return [{"name": y, "values": [v for v in col_nums(table, y) if v is not None], "idx": i,
             "hidden": y in set(spec.get("hidden") or [])} for i, y in enumerate(spec["ys"])]


# ── 새 종류용 데이터 ─────────────────────────────────────────────────────────
TOTAL_NAME = re.compile(r"(합계|총계|소계|누계|시작|기초|기말|期末|잔액|잔고|순이익|순손익|최종|total|subtotal|start|begin|end|net|balance)", re.I)


def waterfall_steps(table, spec, data=None):
    """폭포: 첫 계열 값 = 증감. spec.totals(범주 이름 목록)은 0부터 그린 '합계' 막대(누적을 그 값으로 다시 맞춤).
    spec.show_total(기본: 마지막 범주가 합계가 아니면 참) 이면 끝에 누적 합계 막대를 붙인다.
    반환: [{cat, value, a, b, kind: up|down|total}] — a·b 는 막대 아래·위 끝(누적값)"""
    d = data or category_data(table, spec)
    sers = [s for s in d["series"] if not s["hidden"] and not s.get("role")]
    if not sers:
        return []
    vals = sers[0]["values"]
    totals = set(spec.get("totals") or [])
    out, run = [], 0.0
    for c, v in zip(d["cats"], vals):
        if v is None:
            continue
        if c in totals:
            out.append({"cat": c, "value": v, "a": 0.0, "b": v, "kind": "total"})
            run = v
        else:
            out.append({"cat": c, "value": v, "a": run, "b": run + v, "kind": "up" if v >= 0 else "down"})
            run += v
    show = spec.get("show_total")
    if show is None:
        show = not (out and out[-1]["kind"] == "total")
    if show and out:
        out.append({"cat": spec.get("total_label") or "합계", "value": run, "a": 0.0, "b": run, "kind": "total", "auto": True})
    return out


def bubble_data(table, spec):
    """거품: x·y·크기(spec.size) 숫자 열 + 점 이름(spec.label, 범주 열) + 색 그룹(spec.group). 크기가 0 이하인 점은 뺀다"""
    xs, ys = col_nums(table, spec["x"]), col_nums(table, spec["ys"][0])
    zs = col_nums(table, spec["size"]) if spec.get("size") else [1.0] * len(xs)
    lab = col_values(table, spec["label"]) if spec.get("label") else [""] * len(xs)
    gs = col_values(table, spec["group"]) if spec.get("group") else None
    levels = _levels(gs) if gs else [spec["ys"][0]]
    hidden = set(spec.get("hidden") or [])
    out = [{"name": g, "idx": i, "hidden": g in hidden, "points": []} for i, g in enumerate(levels)]
    for i, (a, b, z) in enumerate(zip(xs, ys, zs)):
        if a is None or b is None or z is None or z <= 0:
            continue
        out[levels.index(gs[i]) if gs else 0]["points"].append((a, b, z, lab[i]))
    return out


_OHLC = {"open": r"(시가|시초가|open(ing)?( price)?|o)", "high": r"(고가|최고가|high( price)?|h)",
         "low": r"(저가|최저가|low( price)?|l)", "close": r"(종가|close|closing( price)?|adj\.? ?close|c)"}
_VOLUME = re.compile(r"^(거래량|거래대금|volume|vol)$", re.I)


def find_ohlc(table):
    """시가·고가·저가·종가(open/high/low/close) 숫자 열 4개를 그 순서로. 없으면 None. 거래량 열은 두 번째 값"""
    nums = [c["name"] for c in table["columns"] if c["type"] == "num"]
    got = {}
    for k, pat in _OHLC.items():
        for n in nums:
            if re.fullmatch(pat, n.strip(), re.I) and n not in got.values():
                got[k] = n
                break
    if len(got) < 4:
        return None, None
    vol = next((n for n in nums if _VOLUME.match(n.strip())), None)
    return [got["open"], got["high"], got["low"], got["close"]], vol


def squarify(items, x, y, w, h):
    """트리맵 배치(Bruls 등의 squarified) — items: [(키, 값>0)] 큰 값 먼저. 반환: [(키, 값, x, y, w, h)]"""
    items = [(k, v) for k, v in items if v and v > 0]
    tot = sum(v for _, v in items)
    if not items or w <= 0 or h <= 0:
        return []
    scale = w * h / tot
    rest = [(k, v * scale, v) for k, v in items]
    out = []

    def worst(row, side):
        s = sum(a for _, a, _ in row)
        mx, mn = max(a for _, a, _ in row), min(a for _, a, _ in row)
        return max(side * side * mx / (s * s), (s * s) / (side * side * mn))
    while rest:
        side = min(w, h)
        row = [rest.pop(0)]
        while rest and worst(row + [rest[0]], side) <= worst(row, side):
            row.append(rest.pop(0))
        s = sum(a for _, a, _ in row)
        if w >= h:  # 왼쪽에 세로 줄
            cw = s / h
            yy = y
            for k, a, v in row:
                hh = a / cw
                out.append((k, v, x, yy, cw, hh))
                yy += hh
            x, w = x + cw, w - cw
        else:  # 위에 가로 줄
            rh = s / w
            xx = x
            for k, a, v in row:
                ww = a / rh
                out.append((k, v, xx, y, ww, rh))
                xx += ww
            y, h = y + rh, h - rh
    return out


def jo(word, with_final, without_final):
    """받침 따라 조사 고르기: jo('실적', '을', '를') → '실적을'. 마지막 글자가 한글이 아니면 받침 없는 쪽"""
    ch = (word or " ").rstrip()[-1:]
    fin = "가" <= ch <= "힣" and (ord(ch) - 0xAC00) % 28
    return word + (with_final if fin else without_final)


def wa(word):
    """받침 따라 '와/과' — 마지막 글자가 한글이 아니면 '와'"""
    ch = (word or " ").rstrip()[-1:]
    return word + ("과" if "가" <= ch <= "힣" and (ord(ch) - 0xAC00) % 28 else "와")


# ── 추천 규칙 ─────────────────────────────────────────────────────────────
def _levels(vals):
    out = []
    for v in vals:
        if v not in out:
            out.append(v)
    return out


def recommend(table, limit=12):
    """데이터 모양을 보고 적합한 그래프 후보 명세 목록. 부적합한 차트는 만들지 않는다."""
    cols = table["columns"]
    rows = table["rows"]
    n = len(rows)
    dates = [c["name"] for c in cols if c["type"] == "date"]
    cats = [c["name"] for c in cols if c["type"] == "cat"]
    nums = [c["name"] for c in cols if c["type"] == "num" and sum(v is not None for v in col_nums(table, c["name"])) >= 2]
    out = []

    def add(kind, why, score, **kw):
        sp = {"kind": kind, "why": why, "score": score, **kw}
        fk = lambda o: repr([o.get(k) for k in fin.FIN_KEYS + ("trend", "trend_n")])  # 주가 지표·추세선 선택이 다르면 다른 후보
        key = (kind, sp.get("x"), tuple(sp.get("ys") or []), sp.get("group"), sp.get("y2"), fk(sp))
        if any((o["kind"], o.get("x"), tuple(o.get("ys") or []), o.get("group"), o.get("y2"), fk(o)) == key for o in out):
            return
        out.append(sp)

    if not nums:
        return []
    by_unit = {}
    for y in nums:
        by_unit.setdefault(unit_of(table, y), []).append(y)
    main_unit, main_ys = max(by_unit.items(), key=lambda kv: len(kv[1]))
    main_ys = main_ys[:8]
    x_time = dates[0] if dates else None
    x_cat = None
    for c in cats:
        lv = _levels(col_values(table, c))
        if 2 <= len(lv) <= 60:
            x_cat = c
            break
    x = x_time or x_cat
    # 긴 형식(x 반복 + 다른 범주 열) → 피벗
    group = None
    if x:
        xs = col_values(table, x)
        nx = len(_levels(xs))
        if nx < n:
            for g in [c for c in cats + dates if c != x]:
                gl = _levels(col_values(table, g))
                pairs = len(set(zip(xs, col_values(table, g))))
                if 2 <= len(gl) <= 8 and pairs >= 0.9 * n:
                    group = g
                    break
    nx = len(_levels(col_values(table, x))) if x else n
    nonneg = lambda ys: all(v is None or v >= 0 for y in ys for v in col_nums(table, y))

    def add_dual(xc, units):
        """단위가 다른 두 묶음: 왼쪽 축 1~2 계열(막대) + 오른쪽 축 1 계열(꺾은선)"""
        if nx > 24:
            return
        lu = main_unit if main_unit in units else units[0]
        ru = next(u for u in units if u != lu)
        left, right = by_unit[lu][:2], by_unit[ru][0]
        add("dual", f"{'·'.join(left)}({lu or '단위 없음'}) 막대 + {right}({ru or '단위 없음'}) 오른쪽 축 꺾은선 — 두 축 눈금이 독립, 추세 비교 시 주의",
            83, x=xc, ys=left, y2=right, dual_shape="bar_line")

    if x and group:
        y0 = main_ys[0]
        k = len(_levels(col_values(table, group)))
        if x == x_time:
            add("line", f"{x} 흐름에 따라 {group}별 {y0} 추이를 비교", 96, x=x, ys=[y0], group=group)
        if nx <= 16:
            add("col_group", f"{x}마다 {group} {k}개를 나란히 비교", 92, x=x, ys=[y0], group=group)
        if nonneg([y0]) and nx <= 24:
            add("col_stack", f"{group}별 구성과 {x}별 합계를 함께", 85, x=x, ys=[y0], group=group)
            add("col_stack100", f"{x}마다 {group} 비중(%)이 어떻게 달라지는지", 72, x=x, ys=[y0], group=group)
        if nx >= 3 and k >= 3:
            add("heatmap", f"{x} × {group} 값을 색 농도로 한눈에", 66, x=x, ys=[y0], group=group)
        add("box", f"{group}별 {y0} 분포(중앙값·사분위) 비교", 50, x=None, ys=[y0], group=group)

    if x_time and main_ys:
        k = len(main_ys)
        add("line", f"{x_time} 순서가 있어 추세를 보기 좋음" + (f" — {k}개 계열 비교" if k > 1 else ""), 95,
            x=x_time, ys=main_ys[:6])
        if nx <= 24:
            add("col", f"{x_time}별 {main_ys[0]} 크기를 막대로", 88 if k == 1 else 70, x=x_time, ys=main_ys[:1])
        if 2 <= k <= 4 and nx <= 12:
            add("col_group", f"{x_time}마다 {k}개 계열을 나란히 비교", 86, x=x_time, ys=main_ys[:4])
        if k >= 2 and nonneg(main_ys):
            add("col_stack", f"계열 합계와 구성을 함께 (같은 단위{('·' + main_unit) if main_unit else ''})", 78, x=x_time, ys=main_ys[:6])
            add("area_stack", "누적 흐름으로 전체 규모 변화와 구성", 64, x=x_time, ys=main_ys[:6])
        if k == 1 and nonneg(main_ys):
            add("area", f"{main_ys[0]} 추이를 면적으로 강조", 60, x=x_time, ys=main_ys[:1])
        if len(by_unit) >= 2:
            units = list(by_unit)[:2]
            a, b = by_unit[units[0]][0], by_unit[units[1]][0]
            add("panels", f"단위가 다른 {a}({units[0] or '단위 없음'})·{b}({units[1] or '단위 없음'}) — 축을 나눠 위아래 두 그래프로", 84,
                x=x_time, ys=[a, b])
            add_dual(x_time, units)
            if all(v and v > 0 for v in col_nums(table, a)[:1] + col_nums(table, b)[:1]):
                add("index", f"단위가 다른 {a}·{b} 증감률을 첫 값=100 으로 맞춰 비교", 74, x=x_time, ys=[a, b])
        elif k >= 2 and nx >= 3:
            firsts = [col_nums(table, y)[0] for y in main_ys]
            if all(f and f > 0 for f in firsts):
                mx = max(max(v for v in col_nums(table, y) if v is not None) for y in main_ys)
                mn = min(max(v for v in col_nums(table, y) if v is not None) for y in main_ys)
                if mn > 0 and mx / mn >= 20 and nx <= 24:  # 같은 단위라도 자릿수가 크게 다르면 큰 쪽 왼축 막대 + 작은 쪽 오른축 선
                    big = max(main_ys, key=lambda y: max(v for v in col_nums(table, y) if v is not None))
                    small = min(main_ys, key=lambda y: max(v for v in col_nums(table, y) if v is not None))
                    add("dual", f"{big}·{small} 크기 자릿수가 달라 {small}을(를) 오른쪽 축 꺾은선으로 — 두 축 눈금이 독립, 추세 비교 시 주의", 66,
                        x=x_time, ys=[big], y2=small, dual_shape="bar_line")
                if mn > 0 and mx / mn >= 5:
                    add("index", "크기 차이가 큰 계열들의 증감률을 첫 값=100 으로 맞춰 비교", 68, x=x_time, ys=main_ys[:4])
        if k >= 3 and nx >= 3:
            add("heatmap", f"{x_time} × 계열 값을 색 농도로", 55, x=x_time, ys=main_ys[:8])

    if x_cat and not group and main_ys and x_cat != x_time:
        k = len(main_ys)
        y0 = main_ys[0]
        long_lbl = max((len(str(v)) for v in col_values(table, x_cat)), default=0) > 6
        if nx <= 30:
            add("bar", f"항목 {nx}개를 크기순으로 — 이름이 길어도 읽기 쉬움", 93 if (long_lbl or nx > 8) else 84,
                x=x_cat, ys=[y0], sort="desc")
        if nx <= 12:
            add("col", f"{x_cat}별 {y0} 비교", 86 if not long_lbl else 70, x=x_cat, ys=[y0])
        if 2 <= k <= 4 and nx <= 12:
            add("col_group" if not long_lbl else "bar_group", f"{x_cat}마다 {k}개 계열을 나란히", 85, x=x_cat, ys=main_ys[:4])
        if k >= 2 and nonneg(main_ys) and nx <= 20:
            add("bar_stack" if long_lbl else "col_stack", "항목별 합계와 구성을 함께", 76, x=x_cat, ys=main_ys[:6], sort="desc")
            add("col_stack100", "항목마다 구성 비율(%) 비교", 70, x=x_cat, ys=main_ys[:6])
        vals = [v for v in col_nums(table, y0) if v is not None]
        if 2 <= nx <= 6 and vals and all(v > 0 for v in vals) and len(vals) == nx:
            add("doughnut", f"{nx}개 항목이 전체에서 차지하는 비중", 80, x=x_cat, ys=[y0])
            add("pie", f"{nx}개 항목의 구성비 — 범주가 적어 원그래프가 적합", 74, x=x_cat, ys=[y0])
        if k >= 3 and nx >= 3:
            add("heatmap", f"{x_cat} × 계열 값을 색 농도로 비교", 62, x=x_cat, ys=main_ys[:8])
        if len(by_unit) >= 2 and nx <= 20:
            units = list(by_unit)[:2]
            add_dual(x_cat, units)
            add("panels", "단위가 다른 두 계열 — 축을 나눠 위아래 두 그래프로", 72, x=x_cat,
                ys=[by_unit[units[0]][0], by_unit[units[1]][0]])

    # ── 새 종류 규칙 ──
    _new_kind_rules(table, add, x_time=x_time, x_cat=x_cat, group=group, nx=nx, main_ys=main_ys, main_unit=main_unit,
                    by_unit=by_unit, nums=nums, cats=cats, n=n, nonneg=nonneg)

    # 숫자 열 쌍 → 산점도 (상관 강한 순 2개)
    pairs = []
    if len(nums) >= 2 and n >= (12 if x_time else 8):
        for i in range(len(nums)):
            for j in range(len(nums)):
                if i < j:
                    lr = linreg(col_nums(table, nums[i]), col_nums(table, nums[j]))
                    if lr:
                        pairs.append((abs(lr["r"]), nums[i], nums[j], lr))
        pairs.sort(reverse=True)
        sc_group = None
        for c in cats:
            lv = _levels(col_values(table, c))
            if 2 <= len(lv) <= 3 and len(lv) < n:
                sc_group = c
                break
        for r, a, b, lr in pairs[:2]:
            add("scatter", f"{wa(a)} {b}의 관계" + (f" — 상관계수 r={lr['r']:.2f}, 추세선 포함" if r >= .3 else " (상관 약함)"),
                (45 if x_time else 70) + 20 * r, x=a, ys=[b], trend=r >= .3)
            if sc_group:
                add("scatter", f"{sc_group}별로 색을 나눠 {a}–{b} 관계 비교", 70 + 5 * r, x=a, ys=[b], group=sc_group, trend=False)

    # 분포 — 산점도의 결과 변수(y)를 먼저
    resp = [p[2] for p in pairs[:2]] if len(nums) >= 2 and n >= (12 if x_time else 8) else []
    if n >= 15:
        dn = [y for y in nums if len(set(v for v in col_nums(table, y) if v is not None)) >= 6]
        dn.sort(key=lambda y: 0 if y in resp else 1)
        if dn:
            add("hist", f"{dn[0]} 값 {n}개가 어느 구간에 몰려 있는지", 75 if not x_time else 52, ys=[dn[0]])
        if len(dn) >= 2:
            add("hist", f"{dn[1]} 값의 분포", 48, ys=[dn[1]])
    bg = next((c for c in cats if 2 <= len(_levels(col_values(table, c))) <= 8 and len(_levels(col_values(table, c))) * 3 <= n), None)
    if bg and not group and n >= 8:
        yb = resp[0] if resp else main_ys[0]
        add("box", f"{bg}별 {yb} 분포(중앙값·사분위·이상값) 비교", 66, x=None, ys=[yb], group=bg)
    if n >= 8 and 2 <= len(main_ys) <= 8 and not group:
        add("box", f"{'·'.join(main_ys[:4])}{'…' if len(main_ys) > 4 else ''} 분포를 중앙값·사분위로 비교", 58 if not x_time else 40,
            ys=main_ys[:8])
    if not out and nums:  # 범주 열이 없고 행이 적음 → 행 번호 기준 막대
        add("col", "행 순서대로 값 비교", 50, x=None, ys=nums[:1])

    # 막대 + 추세선: x 가 기간(순서 있음)이고 막대 후보가 있으면 가장 높은 막대 후보에 직선 추세를 얹은 변형 하나
    if x_time and n >= 4:
        bars = sorted((o for o in out if o["kind"] in ("col", "col_group") and o.get("x") == x_time and not o.get("group")),
                      key=lambda o: -o["score"])
        if bars:
            b0 = bars[0]
            nm = "·".join(b0["ys"][:3]) + ("…" if len(b0["ys"]) > 3 else "")
            add(b0["kind"], f"{nm} 막대 위에 계열별 추세선(최소제곱 직선)을 점선으로 — 들쭉날쭉한 {x_time}별 값의 전체 방향이 보임",
                b0["score"] - 2, x=x_time, ys=list(b0["ys"]), trend="linear")
    ohlc, _vol = find_ohlc(table)
    if ohlc:  # 시가·고가·저가·종가를 쌓거나 합치는 그림은 뜻이 없음
        out[:] = [o for o in out if not (o["kind"] in ("col_stack", "col_stack100", "area_stack", "bar_stack", "heatmap", "col_group", "index")
                                         and len(set(o.get("ys") or []) & set(ohlc)) >= 2)]
    out.sort(key=lambda s: -s["score"])
    out = out[:limit]
    for i, s in enumerate(out):
        s["id"] = f"c{i + 1}"
        s["source"] = "rule"
        s["name"] = KINDS[s["kind"]]
        s.setdefault("title", default_title(table, s))
        s.setdefault("legend", "top")
        s.setdefault("labels", "auto")
    return out


_HINT_RADAR = re.compile(r"(역량|평가|점수|능력|지표|항목별|만족도|score|skill|rating)", re.I)
_HINT_WATER = re.compile(r"(증감|손익|변동|변화|차이|증가|감소|기여|요인|delta|change|profit|loss|variance)", re.I)
_HINT_FUNNEL = re.compile(r"(단계|퍼널|깔때기|전환|방문|유입|가입|구매|결제|지원|합격|서류|면접|stage|funnel|visit|signup|purchase)", re.I)
_HINT_STEP = re.compile(r"(금리|이율|요율|요금|단가|세율|정원|기준|등급|rate|price|tariff|level)", re.I)
_HINT_TARGET = re.compile(r"(목표|계획|기준치|target|goal|plan|budget)", re.I)
_HINT_ACTUAL = re.compile(r"(실적|실제|달성|결과|actual|result)", re.I)
_HINT_LINE = re.compile(r"(목표|계획|예산|target|plan|budget|quota)", re.I)
_ADDITIVE = re.compile(r"(^$|원$|명$|건$|개$|곳$|가구$|대$|톤$|^t$|Wh$|^%$|회$|주$|ha$|㎡$|m2$)")  # 더해서 '전체'가 되는 양
_HINT_PAIR = re.compile(r"(전|후|이전|이후|before|after|남|여|상반기|하반기|목표|실적|계획|예산|결산)", re.I)
_HINT_CYCLE = re.compile(r"^(\d{1,2}\s*월|[1-4]\s*분기|[월화수목금토일](요일)?|[동서남북]{1,3}|N|NE|E|SE|S|SW|W|NW|\d{1,2}\s*시)$", re.I)


def _monotone_down(vals):
    vs = [v for v in vals if v is not None]
    return len(vs) >= 3 and all(a >= b for a, b in zip(vs, vs[1:])) and vs[-1] > 0 and vs[0] / vs[-1] >= 1.5


def _new_kind_rules(table, add, *, x_time, x_cat, group, nx, main_ys, main_unit, by_unit, nums, cats, n, nonneg):
    """방사형·스파크라인·캔들스틱·폭포·깔때기·거품·계단선·막대사탕·덤벨·트리맵·불릿·로즈·누적+선·반원 도넛"""
    k = len(main_ys)
    y0 = main_ys[0] if main_ys else None
    # 캔들스틱: 날짜 + 시가·고가·저가·종가
    ohlc, vol = find_ohlc(table)
    if ohlc and (x_time or x_cat) and n >= 3:
        xx = x_time or x_cat
        vk = {"volume": vol} if vol else {}
        add("candlestick", f"{'·'.join(ohlc)} 열이 있어 기간마다 시작·끝·범위를 캔들로" + (f" — {jo(vol, '은', '는')} 아래 막대로" if vol else ""),
            99, x=xx, ys=ohlc, **vk)
        _fin_rules(add, xx, ohlc, vol, n)
    # 불릿: 실적 + 목표
    tgt = next((c for c in nums if _HINT_TARGET.search(c)), None)
    act = next((c for c in nums if c != tgt and _HINT_ACTUAL.search(c)), None)
    if tgt and not act:
        act = next((c for c in nums if c != tgt and unit_of(table, c) == unit_of(table, tgt)), None)
    xb = x_cat or x_time
    if tgt and act and xb and not group and nx <= 20 and unit_of(table, tgt) == unit_of(table, act):
        add("bullet", f"{jo(act, '을', '를')} 막대로, {jo(tgt, '을', '를')} 가로 눈금으로 — 항목마다 목표 달성 여부가 바로 보임", 95, x=xb, ys=[act], target=tgt)
    # 누적 막대 + 꺾은선: 같은 단위 구성 열 2개 이상 + 목표·평균 같은 기준 열
    line_c = next((c for c in nums if _HINT_LINE.search(c)), None)
    if line_c and xb and not group and nx <= 24:
        parts = [c for c in by_unit.get(unit_of(table, line_c), []) if c != line_c and c != act][:6]
        if len(parts) >= 2 and nonneg(parts) and _ADDITIVE.search(unit_of(table, line_c)):
            add("stack_line", f"{'·'.join(parts[:3])}{'…' if len(parts) > 3 else ''} 누적 막대에 {jo(line_c, '을', '를')} 꺾은선으로 겹쳐 합계와 기준을 비교",
                82, x=xb, ys=parts, y2=line_c)
    if not y0:
        return
    vals0 = [v for v in col_nums(table, y0) if v is not None]
    pos0 = bool(vals0) and all(v > 0 for v in vals0)
    xs_cat = col_values(table, x_cat) if x_cat else []
    # 방사형: 항목(축) 3~12개, 같은 단위 계열 1~5개
    if x_cat and not group and 3 <= nx <= 12 and 1 <= k <= 5 and nonneg(main_ys[:5]) and x_cat != x_time:
        hint = bool(_HINT_RADAR.search(x_cat + " " + " ".join(main_ys) + " " + (table.get("title") or "")))
        if hint or k >= 2:
            add("radar", f"{x_cat} {nx}개를 축으로 놓고 {'·'.join(main_ys[:3])}{'…' if k > 3 else ''}의 모양(강점·약점)을 비교",
                (88 if hint else 58) if k >= 2 else 72, x=x_cat, ys=main_ys[:5])
    # 스파크라인: 시간축 + 계열 3개 이상(또는 그룹 3개 이상) — 계열마다 작은 추세선
    if x_time and nx >= 6:
        if group:
            gl = len(_levels(col_values(table, group)))
            if gl >= 3:
                add("sparkline", f"{group} {gl}개마다 {y0} 추세를 작은 선으로 나란히 — 계열이 많아도 겹치지 않음", 70 + min(gl, 8), x=x_time, ys=[y0], group=group)
        elif len(nums) >= 3:
            ys_sp = nums[:12]
            nu = len({unit_of(table, y) for y in ys_sp})
            add("sparkline", f"계열 {len(ys_sp)}개 추세를 각자 눈금으로 작게 나란히 — " + (f"단위가 {nu}가지라 한 축에 못 담는 흐름을 비교" if nu >= 3 else "계열이 많아도 겹치지 않음"),
                92 if nu >= 3 else 66 + min(len(ys_sp), 10) * 2, x=x_time, ys=ys_sp)
    # 계단선: 값이 가끔만 바뀌는 시계열(금리·요금 등)
    if x_time and not group and nx >= 6:
        for y in main_ys[:2]:
            vs = [v for v in col_nums(table, y) if v is not None]
            same = sum(1 for a, b in zip(vs, vs[1:]) if a == b)
            if len(vs) >= 6 and same >= max(2, 0.3 * (len(vs) - 1)):
                add("step", f"{y} 값이 일정 기간 유지되다 바뀌어 계단 모양으로 — 바뀐 시점이 분명히 보임",
                    97 if _HINT_STEP.search(y) else 80, x=x_time, ys=[y])
                break
    # 폭포: 증감(양·음 섞임) 한 계열
    xw = x_cat or x_time
    if xw and not group and vals0 and any(v < 0 for v in vals0) and any(v > 0 for v in vals0) and 3 <= nx <= 30:
        names = [str(c) for c in _levels(col_values(table, xw))]
        totals = [c for c in names if TOTAL_NAME.search(c)]
        hint = _HINT_WATER.search(y0 + " " + (table.get("title") or "")) or totals
        sp = {"x": xw, "ys": [y0]}
        if totals:
            sp["totals"] = totals
        end = names[-1] if names[-1] in totals else "합계"
        add("waterfall", f"{y0} 증감이 쌓여 {end}에 이르는 과정을 — 늘어난 몫·줄어든 몫을 색으로 구분",
            90 if hint else 66, **sp)
    # 깔때기: 차례대로 줄어드는 단계 값
    if x_cat and not group and 3 <= nx <= 10 and pos0 and len(vals0) == nx and _monotone_down(col_nums(table, y0)):
        hint = _HINT_FUNNEL.search(x_cat + " " + y0 + " " + " ".join(xs_cat) + " " + (table.get("title") or ""))
        add("funnel", f"{x_cat} {nx}단계를 지나며 {jo(y0, '이', '가')} 줄어드는 비율(전환율)을 깔때기로", 94 if hint else 74, x=x_cat, ys=[y0])
    # 덤벨: 같은 단위 두 계열(두 시점·두 집단) 비교
    if xb and not x_time and 3 <= nx <= 20:
        pair = None
        if group:
            gl = _levels(col_values(table, group))
            if len(gl) == 2:
                pair = {"ys": [y0], "group": group}
                nm = f"{gl[0]}·{gl[1]}"
        elif k >= 2:
            two = (main_ys[:2] if k == 2 else ([act, tgt] if tgt and act else None))
            if two:
                pair = {"ys": two}
                nm = "·".join(two)
        if pair:
            hint = _HINT_PAIR.search(nm) or all(is_date(p) for p in nm.split("·"))
            add("dumbbell", f"{x_cat}마다 {nm} 두 값을 점으로 잇고 그 차이를 길이로 비교", 84 if hint else 68, x=x_cat, **pair)
    # 막대사탕·트리맵·로즈·반원 도넛: 범주 + 한 계열
    if x_cat and not group and x_cat != x_time:
        if 8 <= nx <= 40:
            add("lollipop", f"항목 {nx}개를 막대 대신 점과 가는 줄기로 — 많은 항목도 가볍게 크기순 비교", 64, x=x_cat, ys=[y0], sort="desc")
        ids = sum(1 for c in xs_cat if re.fullmatch(r"[A-Za-z]{0,3}[-_ ]?\d+", str(c).strip())) >= 0.6 * max(1, len(xs_cat))
        scores = _HINT_RADAR.search(x_cat + " " + y0 + " " + (table.get("title") or ""))  # 점수·평가는 더해서 '전체'가 안 됨
        if pos0 and 6 <= nx <= 60 and len(vals0) == nx and _ADDITIVE.search(unit_of(table, y0)) and not ids and not scores:
            add("treemap", f"{nx}개 항목이 전체에서 차지하는 몫을 넓이로 — 원그래프로는 많은 항목", 78 if nx > 6 else 68, x=x_cat, ys=[y0])
        if 2 <= nx <= 6 and pos0 and len(vals0) == nx:
            add("semi_doughnut", f"{nx}개 항목 구성비를 반원으로 — 좁은 자리(보고서 요약 칸)에 맞음", 64, x=x_cat, ys=[y0])
    xr = x_cat or x_time
    if xr and not group and 6 <= nx <= 12 and vals0 and all(v >= 0 for v in vals0):
        cyc = [str(c).strip() for c in _levels(col_values(table, xr))]
        if all(_HINT_CYCLE.match(c) for c in cyc):
            add("rose", f"{jo(xr, '이', '가')} 한 바퀴 도는 주기라 원 둘레에 놓고 {y0} 크기를 부채꼴 넓이로", 62, x=xr, ys=[y0])
    # 거품: 숫자 열 3개 — x·y + 크기, 점 이름은 범주 열
    if len(nums) >= 3 and 4 <= n <= 80 and not x_time:
        best = None
        for i, a in enumerate(nums):
            for b in nums[i + 1:]:
                lr = linreg(col_nums(table, a), col_nums(table, b))
                if lr and (best is None or abs(lr["r"]) > best[0]):
                    best = (abs(lr["r"]), a, b)
        if best:
            _, a, b = best
            size = next((c for c in nums if c not in (a, b) and all(v is None or v > 0 for v in col_nums(table, c))), None)
            if size and len({unit_of(table, c) for c in (a, b, size)}) >= 2:
                sp = {"x": a, "ys": [b], "size": size}
                if x_cat:
                    sp["label"] = x_cat
                add("bubble", f"{wa(a)} {b}의 관계에 {jo(size, '을', '를')} 원 크기로 더해 세 값을 한 번에", 76 if n <= 30 else 66, **sp)


def _fin_rules(add, xx, ohlc, vol, n):
    """주가 표(시가·고가·저가·종가) — 이동평균·볼린저·보조지표 칸·매물대 후보 (기간이 지표 창보다 길 때만)"""
    vk = {"volume": vol} if vol else {}
    pv = ["volume"] if vol else []
    if n >= 10:
        add("ohlc", "같은 가격을 OHLC 막대로 — 몸통 없이 세로선(고가~저가) + 왼쪽 눈금 시가·오른쪽 눈금 종가", 60, x=xx, ys=ohlc, **vk)
    ma = [w for w in fin.MA_DEFAULT if w <= n]
    if n >= 20 and len(ma) >= 2:
        add("candlestick", f"캔들에 이동평균 {'·'.join(map(str, ma))}일 선을 겹쳐 단기·장기 추세와 교차를 함께", 98, x=xx, ys=ohlc, ma=ma,
            panels=pv, **vk)
    if n >= 25:
        add("candlestick", "캔들에 볼린저 밴드(20일 평균 ± 2 표준편차)를 겹쳐 가격이 평소 범위를 벗어난 때를 표시", 94, x=xx, ys=ohlc,
            bb={"period": 20, "k": 2}, panels=pv, **vk)
    sub = (["macd"] if n >= 34 else []) + (["rsi"] if n >= 15 else [])
    if sub:
        nm = "·".join(fin.PANEL_KO[p] for p in sub + pv)
        add("candlestick", f"가격 아래에 {nm} 칸을 같은 날짜 축으로 — 추세 강도·과열(RSI 70)·침체(30)를 함께", 96, x=xx, ys=ohlc,
            ma=[w for w in (5, 20) if w <= n], panels=pv + sub, **vk)
    if vol and n >= 20:
        add("candlestick", f"오른쪽에 매물대(가격대별 {vol}) — 거래가 몰린 가격(POC)과 70% 가치 영역", 88, x=xx, ys=ohlc,
            vprofile={"bins": 24, "va": True}, panels=pv, **vk)


def default_title(table, spec):
    ys = spec.get("ys") or []
    k = spec["kind"]
    if k == "scatter":
        return f"{wa(spec['x'])} {ys[0]}의 관계"
    if k in ("hist",):
        return f"{ys[0]} 분포"
    if k == "box":
        return f"{spec['group']}별 {ys[0]} 분포" if spec.get("group") else f"{'·'.join(ys[:3])} 분포"
    if k == "index":
        return f"{'·'.join(ys[:3])} 증감 비교 (첫 값=100)"
    if k == "dual":
        return f"{spec['x']}별 {'·'.join(ys[:2])}과 {spec.get('y2', '')}"
    if trend_mode(spec) and spec.get("x"):
        base = default_title(table, dict(spec, trend=None))
        return f"{base} — {'직선 추세' if trend_mode(spec) == 'linear' else '이동평균 추세'}"
    if k in ("candlestick", "ohlc"):
        extra = fin.describe(spec)
        return f"{spec['x']}별 가격 흐름 " + (f"— {extra}" if extra else ("(OHLC 막대)" if k == "ohlc" else "(시가·고가·저가·종가)"))
    if k == "bubble":
        return f"{wa(spec['x'])} {ys[0]}의 관계 (원 크기: {spec.get('size') or '-'})"
    if k == "bullet":
        return f"{spec['x']}별 {ys[0]}과 {spec.get('target', '목표')}"
    if k == "stack_line":
        return f"{spec['x']}별 {'·'.join(ys[:3])} 구성과 {spec.get('y2', '')}"
    if k == "waterfall":
        return f"{ys[0]} 증감 흐름"
    if k == "funnel":
        return f"{spec['x']}별 {ys[0]} 전환"
    if k == "sparkline":
        return f"{spec['x']}별 추세" + (f" — {spec['group']}별 {ys[0]}" if spec.get("group") else "")
    if k == "dumbbell" and not spec.get("group"):
        return f"{spec['x']}별 {wa(ys[0])} {ys[1] if len(ys) > 1 else ''} 비교"
    if k in ("treemap", "semi_doughnut"):
        return f"{spec['x']}별 {ys[0]} 구성"
    if spec.get("group"):
        return f"{spec['x']}·{spec['group']}별 {ys[0]}"
    if k in ("pie", "doughnut"):
        return f"{spec['x']}별 {ys[0]} 구성비"
    if k == "col_stack100":
        return f"{spec.get('x') or '항목'}별 구성비"
    nm = "·".join(ys[:3]) + ("…" if len(ys) > 3 else "")
    return f"{spec['x']}별 {nm}" if spec.get("x") else nm


def summarize(table, max_cols=12):
    """LLM 에 넘길 요약(숫자는 실제 값만)"""
    out = []
    for c in table["columns"][:max_cols]:
        vals = col_values(table, c["name"])
        if c["type"] == "num":
            ns = [v for v in (num(x) for x in vals) if v is not None]
            if ns:
                out.append(f"- {c['name']} (숫자{', 단위 ' + c['unit'] if c['unit'] else ''}): 최소 {fmt(min(ns))}, 최대 {fmt(max(ns))}, 첫 값 {fmt(ns[0])}, 끝 값 {fmt(ns[-1])}")
        else:
            lv = _levels(vals)
            out.append(f"- {c['name']} ({'날짜·기간' if c['type'] == 'date' else '범주'}, {len(lv)}종): {', '.join(lv[:8])}{' …' if len(lv) > 8 else ''}")
    return f"행 {len(table['rows'])}개" + (f", 표 제목 '{table['title']}'" if table.get("title") else "") + "\n" + "\n".join(out)
